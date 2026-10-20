import csv
import os
import sqlite3
import subprocess
import sys
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook

from conftest import PASSWORD, PIN
from vardiya.backup import create_backup, restore_database
from vardiya.reports import build_report, export_audit, export_report
from vardiya.security import AppError
from vardiya.storage import Database, FileLock, inspect_database


def test_backup_restore_keeps_open_and_retains_original(system, tmp_path):
    s, c, token, eid = system
    sid = s.attendance(eid, PIN, "start", None, str(uuid.uuid4()))
    backup = create_backup(s.db, tmp_path / "backup.sqlite3")
    assert backup.with_suffix(".sqlite3.sha256").exists()
    c.advance(60)
    s.attendance(eid, PIN, "end", sid, str(uuid.uuid4()))
    path = s.db.path
    s.db.close()
    archive = restore_database(backup, path, PASSWORD)
    assert (archive / path.name).exists()
    s.db = Database(path)
    assert s.kiosk_employees()[0]["session_id"] == sid
    assert s.db.meta("clock_alert")
    assert not inspect_database(path)


def test_invalid_backup_and_wrong_password_do_not_modify_target(system, tmp_path):
    s, _, _, _ = system
    backup = create_backup(s.db, tmp_path / "good.sqlite3")
    before = s.db.path.read_bytes()
    with pytest.raises(AppError):
        restore_database(backup, s.db.path, "wrong")
    assert s.db.path.read_bytes() == before
    invalid = tmp_path / "bad.sqlite3"
    invalid.write_bytes(b"not a database")
    with pytest.raises(AppError):
        restore_database(invalid, s.db.path, PASSWORD)
    assert s.db.path.read_bytes() == before


def test_restore_requires_current_password_when_target_is_healthy(system, tmp_path):
    s, _, token, _ = system
    backup = create_backup(s.db, tmp_path / "old.sqlite3")
    new_password = "YeniYonetici!2026"
    s.change_password(token, PASSWORD, new_password)
    path = s.db.path
    s.db.close()
    with pytest.raises(AppError):
        restore_database(backup, path, PASSWORD)
    restore_database(backup, path, new_password)
    s.db = Database(path)
    assert s.login_admin(PASSWORD)  # Credentials return to the backup snapshot.


def test_restore_corrupt_target_using_backup_admin(system, tmp_path):
    s, _, _, _ = system
    backup = create_backup(s.db, tmp_path / "good.sqlite3")
    path = s.db.path
    s.db.close()
    path.write_bytes(b"damaged original preserved")
    archive = restore_database(backup, path, PASSWORD)
    assert (archive / path.name).read_bytes() == b"damaged original preserved"
    s.db = Database(path)
    assert not inspect_database(path)


def test_checksum_rejects_modified_backup(system, tmp_path):
    s, _, _, _ = system
    backup = create_backup(s.db, tmp_path / "good.sqlite3")
    backup.with_suffix(".sqlite3.sha256").write_text("0" * 64, encoding="utf-8")
    with pytest.raises(AppError, match="SHA"):
        restore_database(backup, s.db.path, PASSWORD)


def test_restart_recovers_hot_rollback_journal(system):
    s, _, _, eid = system
    path = s.db.path
    s.db.close()
    script = """
import os, sqlite3, sys
c = sqlite3.connect(sys.argv[1], isolation_level=None)
c.execute('PRAGMA cache_size=1')
c.execute('PRAGMA synchronous=FULL')
c.execute('BEGIN IMMEDIATE')
for i in range(100):
    c.execute("INSERT INTO meta VALUES(?,?)", ('crash-'+str(i), 'x'*8000))
c.execute("UPDATE employees SET name='UNCOMMITTED'")
os._exit(17)
"""
    result = subprocess.run([sys.executable, "-c", script, str(path)], capture_output=True)
    assert result.returncode == 17
    s.db = Database(path)
    assert s.db.meta("crash-0") is None
    assert s.kiosk_employees()[0]["name"] == "Ayşe Yılmaz"
    assert not inspect_database(path)


def test_direct_record_tampering_is_reported(system):
    s, _, _, eid = system
    s.db.conn.execute("UPDATE employees SET rate_cents=999 WHERE id=?", (eid,))
    assert any("uyuşmuyor" in error for error in inspect_database(s.db.path))


def test_backup_failure_leaves_no_final_file(system, tmp_path, monkeypatch):
    s, _, _, _ = system
    import vardiya.backup as module
    monkeypatch.setattr(module, "inspect_database", lambda _: ["simulated failure"])
    path = tmp_path / "failed.sqlite3"
    with pytest.raises(AppError):
        create_backup(s.db, path)
    assert not path.exists()
    assert not list(tmp_path.glob(".backup-*"))


def test_refuse_backup_overwrite(system, tmp_path):
    s, _, _, _ = system
    path = create_backup(s.db, tmp_path / "backup.sqlite3")
    with pytest.raises(AppError):
        create_backup(s.db, path)
    with pytest.raises(AppError):
        create_backup(s.db, s.db.path)


def test_process_crash_rolls_back_uncommitted_and_preserves_committed(system):
    s, c, _, eid = system
    s.attendance(eid, PIN, "start", None, str(uuid.uuid4()))
    script = """
import os, sqlite3, sys
c = sqlite3.connect(sys.argv[1], isolation_level=None)
c.execute('PRAGMA synchronous=FULL')
c.execute('BEGIN IMMEDIATE')
c.execute("UPDATE employees SET name='UNCOMMITTED' WHERE id=?", (int(sys.argv[2]),))
os._exit(17)
"""
    child = subprocess.run([sys.executable, "-c", script, str(s.db.path), str(eid)], capture_output=True)
    assert child.returncode == 17
    assert s.kiosk_employees()[0]["name"] == "Ayşe Yılmaz"
    assert s.kiosk_employees()[0]["session_id"] is not None
    assert not inspect_database(s.db.path)


def test_process_lock_released_on_exit(tmp_path):
    lock = FileLock(tmp_path)
    script = """
from pathlib import Path
import sys
from vardiya.storage import FileLock
from vardiya.security import AppError
try:
    lock = FileLock(Path(sys.argv[1]))
except AppError:
    sys.exit(19)
lock.close()
"""
    try:
        child = subprocess.run([sys.executable, "-c", script, str(tmp_path)], capture_output=True)
        assert child.returncode == 19
    finally:
        lock.close()
    child = subprocess.run([sys.executable, "-c", script, str(tmp_path)], capture_output=True)
    assert child.returncode == 0, child.stderr


def test_exports_totals_unicode_and_formula_injection(system, tmp_path):
    s, c, token, eid = system
    s.update_employee(token, eid, '=HYPERLINK("x")', 15000, True, "Dışa aktarım testi")
    s.manual_session(token, eid, c.now - 3600, c.now, 15000, "Eksik kayıt tamamlandı", str(uuid.uuid4()))
    s.add_payment(token, eid, c.now, 3000, "Avans", "@SUM(A1:A2)", str(uuid.uuid4()))
    r = build_report(s, token, date(2026, 10, 20), date(2026, 10, 20))
    csv_path, xlsx_path = tmp_path / "report.csv", tmp_path / "report.xlsx"
    export_report(s, token, r, csv_path)
    export_report(s, token, r, xlsx_path)
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert rows[3][1].startswith("'=")
    assert rows[3][3] == "150.00"
    assert rows[3][-1] == "120.00"
    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == ["Özet", "Günlük", "Giriş çıkış", "Ödeme avans", "Açıklama"]
    assert workbook["Özet"]["D2"].value == 150
    assert workbook["Özet"]["J2"].value == 120
    assert workbook["Ödeme avans"]["G2"].value.startswith("'@")
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    export_audit(s, token, tmp_path / "audit.csv")
    with pytest.raises(AppError):
        export_audit(s, token, s.db.path)


def test_failed_export_preserves_existing_file(system, tmp_path, monkeypatch):
    s, _, token, _ = system
    r = build_report(s, token, date(2026, 10, 20), date(2026, 10, 20))
    path = tmp_path / "report.xlsx"
    path.write_bytes(b"old report")
    import openpyxl
    def fail(*args):
        raise OSError("disk full")
    monkeypatch.setattr(openpyxl.Workbook, "save", fail)
    with pytest.raises(OSError):
        export_report(s, token, r, path)
    assert path.read_bytes() == b"old report"


def test_negative_csv_balance_is_numeric(system, tmp_path):
    s, c, token, eid = system
    s.add_payment(token, eid, c.now, 3000, "Avans", "Çalışma öncesi avans", str(uuid.uuid4()))
    report = build_report(s, token, date(2026, 10, 20), date(2026, 10, 20))
    path = tmp_path / "negative.csv"
    export_report(s, token, report, path)
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert rows[3][-1] == "-30.00"


def test_unsupported_schema_is_not_silently_replaced(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE important(value TEXT)")
    conn.execute("INSERT INTO important VALUES('keep')")
    conn.commit()
    conn.close()
    before = path.read_bytes()
    with pytest.raises(AppError):
        Database(path)
    assert path.read_bytes() == before


def test_maintenance_output_survives_legacy_console_encoding(tmp_path):
    """Redirected output on a cp1252 console must not crash on Turkish text."""
    env = {**os.environ, "PYTHONIOENCODING": "cp1252", "PYTHONUTF8": "0"}
    child = subprocess.run([sys.executable, "-m", "vardiya", "--data-dir", str(tmp_path / "bos"), "--check"],
                           capture_output=True, env=env)
    assert child.returncode == 1, child.stderr
    assert b"Traceback" not in child.stderr
    assert "okunamadı" in child.stdout.decode("utf-8")
