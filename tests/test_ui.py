import time
import uuid
from datetime import date

from PySide6.QtCore import QDate, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QDialog, QLabel

from conftest import PASSWORD, PIN, ts
from vardiya.ui import AdminDialog, EmployeeDialog, MainWindow, PaymentDialog, PinDialog, SessionDialog
from vardiya.widgets import PersonRing, WelcomeDialog


def ring(window, index=0):
    widget = window.grid.itemAt(index).widget()
    assert isinstance(widget, PersonRing)
    return widget


def test_kiosk_touch_pin_start_finish_and_cancel(qapp, system):
    s, c, _, _ = system
    window = MainWindow(s)
    window.show()
    qapp.processEvents()
    assert window.grid.count() == 1
    assert ring(window).state == "out" and "Dışarıda" in ring(window).text()
    employee = s.kiosk_employees()[0]
    dialog = PinDialog(s, employee, window)
    dialog.show()
    for digit in PIN:
        dialog.press(digit)
    QTest.mouseClick(dialog.confirm_button, Qt.MouseButton.LeftButton)
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert dialog.result_sid and dialog.pin.text() == ""
    window.refresh()
    assert ring(window).state == "in" and "İçeride" in ring(window).text()
    c.advance(60)
    dialog = PinDialog(s, s.kiosk_employees()[0], window)
    dialog.pin.setText(PIN)
    dialog.confirm()
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert s.kiosk_employees()[0]["session_id"] is None
    cancelled = PinDialog(s, s.kiosk_employees()[0], window)
    cancelled.pin.setText(PIN)
    cancelled.reject()
    assert cancelled.pin.text() == ""
    window.close()


def test_wrong_pin_stays_in_dialog(qapp, system):
    s, _, _, _ = system
    dialog = PinDialog(s, s.kiosk_employees()[0])
    dialog.pin.setText("0192")
    dialog.confirm()
    assert dialog.result() != QDialog.DialogCode.Accepted
    assert dialog.error.text()
    assert s.kiosk_employees()[0]["session_id"] is None
    dialog.reject()


def test_welcome_screen_after_start_and_finish(qapp, system):
    s, c, _, eid = system
    window = MainWindow(s)
    dialog = PinDialog(s, s.kiosk_employees()[0], window)
    dialog.pin.setText(PIN)
    dialog.confirm()
    summary = s.staff_summary(eid, dialog.result_sid)
    assert summary["session_seconds"] == 0
    welcome = WelcomeDialog("ay", "in", "hoş geldin, ayşe", "giriş 12:00", "", window, timeout_ms=50)
    welcome.show()
    QTest.qWait(200)
    assert welcome.result() == QDialog.DialogCode.Accepted
    c.advance(3 * 3600)
    dialog = PinDialog(s, s.kiosk_employees()[0], window)
    dialog.pin.setText(PIN)
    dialog.confirm()
    summary = s.staff_summary(eid, dialog.result_sid)
    assert summary["session_seconds"] == 3 * 3600 and summary["month_seconds"] == 3 * 3600
    window.close()


def test_admin_pages_forms_and_inactivity(qapp, system):
    s, _, token, _ = system
    panel = AdminDialog(s, token)
    panel.show()
    qapp.processEvents()
    assert panel.pages.count() == 6 and panel.nav.count() == 6
    assert panel.people_table.rowCount() == 1
    assert panel.report_table.rowCount() == 1
    assert panel.tile_inside.value.text() == "0"
    for index in range(panel.nav.count()):
        panel.nav.setCurrentRow(index)
        qapp.processEvents()
        assert panel.pages.currentIndex() == index
    for form in (EmployeeDialog(s, token), SessionDialog(s, token), PaymentDialog(s, token)):
        form.show()
        qapp.processEvents()
        form.reject()
    panel.last_input = time.monotonic() - 301
    panel.check_timeout()
    assert not panel.isVisible()
    assert token not in s._tokens


def test_search_filters_rings(qapp, system):
    s, _, _, _ = system
    window = MainWindow(s)
    window.search.setText("bulunamaz")
    qapp.processEvents()
    assert "bulunamadı" in window.grid.itemAt(0).widget().text()
    window.search.setText("AYŞE")
    qapp.processEvents()
    assert "Ayşe" in ring(window).text()
    window.close()


def test_clock_alert_keeps_kiosk_usable_and_flags(qapp, system):
    s, c, token, _ = system
    window = MainWindow(s)
    window.show()
    s.tick()
    c.now -= 600
    window.heartbeat()
    assert window.cards.isEnabled()
    assert "işaretleniyor" in window.warning.text()
    dialog = PinDialog(s, s.kiosk_employees()[0], window)
    dialog.pin.setText(PIN)
    dialog.confirm()
    assert dialog.result() == QDialog.DialogCode.Accepted
    token = s.login_admin(PASSWORD)
    assert s.pending_reviews(token) == 1
    panel = AdminDialog(s, token)
    assert panel.review_note.text()
    assert panel.tile_review.value.text() == "1"
    panel.session_filter.review.setChecked(True)
    panel.refresh()
    assert panel.session_table.rowCount() == 1
    assert panel.session_table.item(0, 6).text()
    panel.reject()
    window.close()


def test_locked_staff_ring_and_admin_status(qapp, system):
    s, c, token, eid = system
    from vardiya.security import AppError
    for _ in range(5):
        try:
            s.attendance(eid, "1928", "start", None, str(uuid.uuid4()))
        except AppError:
            pass
    window = MainWindow(s)
    window.refresh()
    assert ring(window).state == "locked" and "Kilitli" in ring(window).text()
    panel = AdminDialog(s, token)
    assert "Kilitli" in panel.people_table.item(0, 3).text()
    s.unlock_employee(token, eid, "Test için kilit kaldırıldı")
    panel.refresh()
    assert "Kilitli" not in panel.people_table.item(0, 3).text()
    window.heartbeat()
    assert ring(window).state == "out"
    panel.reject()
    window.close()


def test_admin_shows_turkish_money_actions_and_filters(qapp, system):
    s, c, token, eid = system
    s.manual_session(token, eid, c.now - 40 * 86400, c.now - 40 * 86400 + 3600, 15000,
                     "Eski kayıt eklendi", str(uuid.uuid4()))
    s.manual_session(token, eid, c.now - 3600, c.now, 15000, "Yeni kayıt eklendi", str(uuid.uuid4()))
    panel = AdminDialog(s, token)
    assert panel.people_table.item(0, 2).text() == "150,00"
    assert not panel.people_table.item(0, 1).icon().isNull()
    assert panel.session_table.rowCount() == 1  # default: this month
    assert panel.audit_table.item(0, 3).text() == "Kayıt eklendi"
    assert panel.audit_table.item(0, 2).text() == "Yönetici"
    assert panel.audit_table.item(0, 4).text() == f"Kayıt {panel.session_rows[0]['id']}"
    assert len(panel.session_table.item(0, 2).text()) == len("20.10.2026 11:00")  # no seconds on screen
    assert "MB" in panel.backup_note.text()
    assert panel.tile_hours.value.text() == "1 sa"
    assert panel.timeline.rows and panel.timeline.rows[0][3]
    new = s.add_employee(token, "Can Arslan", "5738", 12000)
    panel.refresh()
    assert panel.session_filter.employee.findData(new) > 0
    panel.reject()


def test_forgotten_exit_ring_pin_warning_and_report(qapp, system):
    s, c, _, eid = system
    c.advance(ts("2026-10-20T17:00") - c.now)
    sid = s.attendance(eid, PIN, "start", None, str(uuid.uuid4()))
    window = MainWindow(s)
    window.show()
    window.refresh()
    assert ring(window).state == "in"
    c.advance(ts("2026-10-21T05:10") - c.now)
    window.heartbeat()
    assert ring(window).state == "forgotten" and "Çıkış unutuldu" in ring(window).text()
    assert window.warning.text() == "Yönetici için 1 not var."
    dialog = PinDialog(s, s.kiosk_employees()[0], window)
    assert dialog.forgotten
    assert any("önceki iş gününden" in w.text() for w in dialog.findChildren(QLabel))
    dialog.pin.setText(PIN)
    dialog.confirm()
    assert dialog.result() == QDialog.DialogCode.Accepted
    token = s.login_admin(PASSWORD)
    panel = AdminDialog(s, token)
    assert panel.report_filter.range[1].isoformat() == "2026-10-21"
    assert any(panel.session_table.item(r, 0).text() == str(sid) and panel.session_table.item(r, 6).text()
               for r in range(panel.session_table.rowCount()))
    assert "05:00" in panel.day_note.text()
    panel.reject()
    window.close()


def test_report_today_uses_business_day(qapp, system):
    s, c, token, _ = system
    c.advance(ts("2026-10-21T00:30") - c.now)
    token = s.login_admin(PASSWORD)
    panel = AdminDialog(s, token)
    panel.report_filter.chips["today"].click()
    assert panel.report_filter.range[0].isoformat() == "2026-10-20"
    assert panel.report["first"] == panel.report["last"] == "2026-10-20"
    assert panel.report_table.columnCount() == 8
    assert "salı" in panel.today_subtitle.text()
    panel.reject()


def test_theme_switch_persists_and_applies(qapp, system):
    from PySide6.QtGui import QPalette
    from vardiya.widgets import current
    s, _, token, _ = system
    window = MainWindow(s)
    panel = AdminDialog(s, token, window)
    light_chip = next(chip for chip, name in panel.theme_chips if name == "light")
    light_chip.click()
    assert s.theme == "light" and current()["name"] == "light"
    assert qapp.palette().color(QPalette.ColorRole.Window).lightness() > 200
    assert light_chip.isChecked()
    night_chip = next(chip for chip, name in panel.theme_chips if name == "night")
    night_chip.click()
    assert s.theme == "night" and current()["name"] == "night"
    assert qapp.palette().color(QPalette.ColorRole.Window).lightness() < 30
    assert night_chip.isChecked() and not light_chip.isChecked()
    dark_chip = next(chip for chip, name in panel.theme_chips if name == "dark")
    dark_chip.click()
    assert s.theme == "dark" and current()["name"] == "dark"
    assert qapp.palette().color(QPalette.ColorRole.Window).lightness() < 60
    assert any(a["action"] == "SETTING_CHANGE" for a in s.audit_rows(token))
    panel.reject()
    window.close()


def test_logo_and_fonts_are_bundled(qapp):
    from PySide6.QtGui import QFontDatabase
    from vardiya.widgets import ASSETS, app_icon, load_fonts, logo_pixmap
    load_fonts()
    families = set(QFontDatabase.families())
    assert "Jost" in families and "Quicksand" in families
    assert not logo_pixmap("#f6ece3", 80).isNull()
    assert not app_icon().isNull()
    assert (ASSETS / "icon.ico").exists() and (ASSETS / "fonts" / "OFL-Jost.txt").exists()


def test_pin_dots_follow_typed_digits(qapp, system):
    s, _, _, _ = system
    dialog = PinDialog(s, s.kiosk_employees()[0])
    for digit in "482":
        dialog.press(digit)
    assert dialog.dots.filled == 3
    dialog.press("Sil")
    assert dialog.dots.filled == 2
    dialog.pin.setText("0009")
    dialog.confirm()
    assert dialog.dots.filled == 0 and dialog.dots.error
    dialog.reject()


def test_touch_animations_finish_and_leave_nothing_running(qapp, system):
    from PySide6.QtCore import QAbstractAnimation
    from vardiya.widgets import PinDots, RingWidget
    s, _, _, _ = system
    window = MainWindow(s)
    window.show()
    qapp.processEvents()
    tile = ring(window)
    tile.pressed.emit()
    QTest.qWait(150)
    assert tile.press_level > 0.8           # squeezed while the finger is down
    tile.released.emit()
    QTest.qWait(450)
    assert abs(tile.press_level) < 0.01     # sprang back to rest
    dots = PinDots()
    dots.set_filled(1)
    QTest.qWait(250)
    assert dots.pop == 1.0
    dots.set_filled(0, error=True)
    QTest.qWait(450)
    assert dots.shake == 1.0
    big = RingWidget(60)
    big.set_ring("in", 1.0, "ca")
    big.draw_in(ms=200, delay=0)
    QTest.qWait(350)
    assert big.reveal == 1.0
    for owner in (tile, dots, big):
        for name in ("press", "hover", "pop", "shake", "reveal"):
            anim = getattr(owner, "_anim_" + name, None)
            assert anim is None or anim.state() == QAbstractAnimation.State.Stopped
    window.close()


def test_animations_can_be_switched_off(qapp, monkeypatch):
    from vardiya import widgets
    monkeypatch.setattr(widgets, "ANIMATIONS", False)
    dots = widgets.PinDots()
    dots.set_filled(2)
    assert dots.pop == 1.0                  # applied instantly, no timer created
    big = widgets.RingWidget(60)
    big.draw_in()
    assert big.reveal == 1.0 and getattr(big, "_anim_reveal", None) is None


def test_closed_main_window_stops_its_timers(qapp, system):
    s, _, _, _ = system
    window = MainWindow(s)
    window.show()
    assert window.timer.isActive() and window.clock_timer.isActive()
    window.close()
    assert not window.timer.isActive() and not window.clock_timer.isActive()


def test_clock_and_greeting_visible_immediately(qapp, system):
    s, _, _, _ = system
    window = MainWindow(s)
    window.show()
    assert window.clock_label.text() == "12:00"
    assert "salı" in window.greeting.text()
    window.close()
    window.show()
    assert window.timer.isActive() and window.clock_timer.isActive()


def test_pin_submits_itself_after_four_digits(qapp, system):
    s, _, _, _ = system
    dialog = PinDialog(s, s.kiosk_employees()[0])
    dialog.show()
    for digit in PIN:
        dialog.press(digit)
    assert dialog.dots.filled == 4
    QTest.qWait(400)
    assert dialog.result() == QDialog.DialogCode.Accepted and dialog.result_sid
    dialog.confirm()                     # a late extra tap does nothing
    assert len(s.sessions(s.login_admin(PASSWORD))) == 1


def test_admin_login_and_setup_use_the_keypad(qapp, tmp_path):
    from conftest import FakeClock
    from vardiya.service import Service
    from vardiya.storage import Database
    from vardiya.ui import admin_login_dialog, setup_dialog
    db = Database(tmp_path / "v.sqlite3")
    s = Service(db, FakeClock(ts("2026-10-20T12:00")))
    dlg = setup_dialog(s)
    for d in "7391":
        dlg.press(d)
    dlg.confirm()
    assert "tekrar" in dlg.prompt.text()
    for d in "7392":
        dlg.press(d)
    dlg.confirm()                                    # mismatch: starts over, nothing saved
    assert not s.configured and not dlg.error.isHidden()
    for code in ("7391", "7391"):
        for d in code:
            dlg.press(d)
        dlg.confirm()
    assert s.configured and dlg.result() == QDialog.DialogCode.Accepted
    login = admin_login_dialog(s)
    login.show()
    for d in "7391":
        login.press(d)
    QTest.qWait(400)                                 # submits itself on the 4th digit
    assert login.result() == QDialog.DialogCode.Accepted and login.result_value in s._tokens
    db.close()


def test_keypad_never_focuses_a_text_field(qapp, system):
    s, _, _, _ = system
    dialog = PinDialog(s, s.kiosk_employees()[0])
    dialog.show()
    qapp.processEvents()
    assert dialog.pin.focusPolicy() == Qt.FocusPolicy.NoFocus
    assert not dialog.pin.testAttribute(Qt.WidgetAttribute.WA_InputMethodEnabled)
    QTest.keyClicks(dialog, "48")                    # a physical keyboard still works
    assert dialog.pin.text() == "48"
    QTest.keyClick(dialog, Qt.Key.Key_Backspace)
    assert dialog.pin.text() == "4"
    dialog.reject()


def test_day_by_day_navigation(qapp, system):
    s, _, token, _ = system
    panel = AdminDialog(s, token)
    bar = panel.session_filter
    assert bar.preset == "this_month" and bar.chips["this_month"].isChecked()
    assert bar.caption.text() == "ekim 2026 · 20 gün"
    assert not bar.next.isEnabled()                     # nothing after today
    bar.prev.click()                                    # this month → yesterday, not last month
    assert bar.range == (date(2026, 10, 19),) * 2 and bar.chips["yesterday"].isChecked()
    bar.prev.click()                                    # the day before: no chip is lit
    assert bar.range == (date(2026, 10, 18),) * 2
    assert not any(chip.isChecked() for chip in bar.chips.values())
    assert bar.apply_button.isHidden()
    bar.next.click()
    bar.next.click()
    assert bar.range == (date(2026, 10, 20),) * 2 and bar.chips["today"].isChecked()
    bar.prev.click()
    # The period is shared: the earnings and payments pages follow.
    assert panel.report["first"] == panel.report["last"] == "2026-10-19"
    assert panel.payment_filter.range == bar.range and panel.payment_filter.chips["yesterday"].isChecked()
    panel.report_filter.chips["last_month"].click()
    assert bar.range == (date(2026, 9, 1), date(2026, 9, 30))
    assert panel.report["first"] == "2026-09-01" and panel.report["last"] == "2026-09-30"
    panel.reject()


def test_custom_range_waits_for_filter_button(qapp, system):
    s, c, token, eid = system
    s.manual_session(token, eid, c.now - 10 * 86400, c.now - 10 * 86400 + 3600, 15000,
                     "Eski kayıt eklendi", str(uuid.uuid4()))
    panel = AdminDialog(s, token)
    panel.show()
    bar = panel.session_filter
    assert panel.session_table.rowCount() == 1
    bar.chips["custom"].click()
    assert bar.chips["custom"].isChecked() and not bar.apply_button.isHidden()
    bar.first.setDate(QDate(2026, 10, 12))
    bar.last.setDate(QDate(2026, 10, 15))
    QTest.qWait(400)
    assert panel.session_table.rowCount() == 1            # typed dates alone change nothing
    assert bar.range == (date(2026, 10, 1), date(2026, 10, 20))
    assert bar.caption.isHidden()                         # the old caption would be wrong now
    panel.refresh()                                       # nor does an unrelated refresh
    assert panel.session_table.rowCount() == 1
    bar.apply_button.click()
    assert bar.range == (date(2026, 10, 12), date(2026, 10, 15)) and bar.caption.text() == "4 gün"
    assert not bar.caption.isHidden()
    assert panel.session_table.rowCount() == 0
    assert panel.report["first"] == "2026-10-12" and panel.payment_filter.range == bar.range
    bar.first.setDate(QDate(2026, 10, 18))                # start past the end: the end follows
    assert bar.last.date() == QDate(2026, 10, 18)
    bar.last.setDate(QDate(2026, 10, 5))                  # end before the start: the start follows
    assert bar.first.date() == QDate(2026, 10, 5)
    assert bar.last.maximumDate() == QDate(2026, 10, 20)  # no future dates
    bar.apply_button.click()
    assert panel.session_table.rowCount() == 0
    bar.first.setDate(QDate(2026, 10, 10))
    bar.last.setDate(QDate(2026, 10, 10))
    bar.apply_button.click()
    assert panel.session_table.rowCount() == 1
    panel.reject()


def test_report_filters_by_employee_and_export_uses_applied_period(qapp, system, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    s, c, token, eid = system
    other = s.add_employee(token, "Can Arslan", "5738", 12000)
    s.manual_session(token, other, c.now - 3600, c.now, 12000, "Kayıt eklendi", str(uuid.uuid4()))
    panel = AdminDialog(s, token)
    panel.show()
    bar = panel.report_filter
    assert panel.report_table.rowCount() == 2
    bar.employee.setCurrentIndex(bar.employee.findData(other))
    QTest.qWait(400)
    assert panel.report_table.rowCount() == 1 and panel.report_table.item(0, 0).text() == "Can Arslan"
    bar.first.setDate(QDate(2026, 9, 1))                  # typed, not applied
    target = tmp_path / "rapor.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    panel.export(".xlsx")
    sheet = load_workbook(target)["Özet"]
    assert [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["Can Arslan"]
    notes = dict(load_workbook(target)["Açıklama"].iter_rows(min_row=2, values_only=True))
    assert notes["Dönem"] == "2026-10-01 — 2026-10-20"
    panel.reject()


def test_windows_fit_small_screens_and_inside_count(qapp, system):
    s, _, token, eid = system
    window = MainWindow(s)
    window.show()
    assert "kimse" in window.inside_label.text()
    s.attendance(eid, PIN, "start", None, str(uuid.uuid4()))
    window.refresh()
    assert window.inside_label.text() == "içeride 1 kişi"
    area = window.screen().availableGeometry()
    panel = AdminDialog(s, token, window)
    assert panel.width() <= area.width() and panel.height() <= area.height()
    pin = PinDialog(s, s.kiosk_employees()[0], window)
    pin.show()
    qapp.processEvents()
    assert pin.sizeHint().height() <= 720
    pin.reject()
    panel.reject()
    window.close()


def test_filters_apply_by_themselves(qapp, system):
    s, c, token, eid = system
    s.manual_session(token, eid, c.now - 3 * 86400, c.now - 3 * 86400 + 3600, 15000,
                     "Eski kayıt eklendi", str(uuid.uuid4()))
    panel = AdminDialog(s, token)
    panel.show()                                     # hidden panels never auto-refresh
    assert panel.session_table.rowCount() == 1
    panel.session_filter.prev.click()                # 19 Ekim: boş, at once
    assert panel.session_table.rowCount() == 0
    panel.session_filter.chips["this_month"].click()
    assert panel.session_table.rowCount() == 1
    panel.session_filter.review.setChecked(True)     # the checkbox applies after a short debounce
    QTest.qWait(400)
    assert panel.session_table.rowCount() == 0
    panel.reject()


def test_closing_admin_with_pending_filter_does_not_refresh(qapp, system):
    s, _, token, _ = system
    panel = AdminDialog(s, token)
    panel.show()
    panel.session_filter.review.setChecked(True)     # starts the 200 ms auto-apply timer
    panel.reject()                                    # closes before it fires
    assert not panel.session_filter._apply.isActive()
    QTest.qWait(400)                                  # no "oturum kapandı" box, no crash


def test_double_click_opens_editor_and_audit_shows_names(qapp, system, monkeypatch):
    s, c, token, eid = system
    s.attendance(eid, PIN, "start", None, str(uuid.uuid4()))
    opened = []
    monkeypatch.setattr(EmployeeDialog, "exec", lambda self: opened.append("employee") or 0)
    monkeypatch.setattr(SessionDialog, "exec", lambda self: opened.append("session") or 0)
    panel = AdminDialog(s, token)
    for table in (panel.people_table, panel.session_table):
        table.setCurrentCell(0, 1)
        table.cellDoubleClicked.emit(0, 1)
    assert opened == ["employee", "session"]
    start = next(r for r in range(panel.audit_table.rowCount()) if panel.audit_table.item(r, 3).text() == "İşe başladı")
    assert panel.audit_table.item(start, 2).text() == "Ayşe Yılmaz"
    panel.reject()


def test_audit_party_words():
    from vardiya.ui import audit_party
    names = {5: "Can Arslan"}
    assert audit_party("staff:5", names) == "Can Arslan"
    assert audit_party("employee:9", names) == "Personel 9"
    assert audit_party("payment:3", names) == "Ödeme 3"
    assert audit_party("system", names) == "Sistem"
    assert audit_party("day_start_hour", names) == "İş günü saati"
    assert audit_party("unknown", names) == "unknown"


def test_health_problems_reach_kiosk_count_and_admin_page(qapp, system, monkeypatch):
    s, c, token, _ = system
    window = MainWindow(s)
    window.show()
    window.heartbeat()                       # takes today's automatic backup first
    assert not window.warning.isVisible()
    monkeypatch.setattr(s.health, "problems", lambda: ["Diskte yalnızca 100 MB boş yer kaldı; yedekler alınamayabilir."])
    window.heartbeat()
    assert window.warning.text() == "Yönetici için 1 not var."
    panel = AdminDialog(s, token)
    assert "100 MB" in panel.health_note.text() and not panel.health_note.isHidden()
    panel.reject()
    window.close()
