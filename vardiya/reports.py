import csv
import os
import tempfile
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

from .security import AppError


def round_cents(value: Fraction) -> int:
    """Nonnegative, exact half-up rounding, once per employee/day/status."""
    return (2 * value.numerator + value.denominator) // (2 * value.denominator)


def midnight(day: date, zone, hour: int = 0) -> int:
    """Start of business day `day`: local `hour`:00 (0 = calendar midnight)."""
    return int(datetime.combine(day, time(hour=hour), zone).timestamp())


def business_day(timestamp: int, zone, hour: int = 0) -> date:
    """Business day a moment belongs to; before `hour`:00 counts as the previous day."""
    return (datetime.fromtimestamp(timestamp, zone) - timedelta(hours=hour)).date()


def split_days(start: int, end: int, zone, hour: int = 0):
    """Split [start, end) at business-day boundaries (local `hour`:00)."""
    while start < end:
        day = business_day(start, zone, hour)
        stop = min(end, midnight(day + timedelta(days=1), zone, hour))
        yield day.isoformat(), stop - start
        start = stop


def build_report(service, token, first: date, last: date):
    service.require_admin(token)
    if last < first or (last - first).days > 365:
        raise AppError("Rapor aralığı 1–366 gün olmalı.")
    zone, now, hour = service.zone, service.clock.wall(), service.day_start_hour
    begin, stop = midnight(first, zone, hour), midnight(last + timedelta(days=1), zone, hour)
    # One consistent DB snapshot for detail, payments and totals.
    with service.db.transaction():
        employees = [dict(r) for r in service.db.conn.execute("SELECT * FROM employees ORDER BY name,id")]
        sessions = [dict(r) for r in service.db.conn.execute(
            "SELECT * FROM sessions WHERE voided=0 AND start_at<?", (stop,))]
        payments = [dict(r) for r in service.db.conn.execute(
            "SELECT * FROM payments WHERE voided=0 AND paid_at<?", (stop,))]
        clock_alert = service.db.meta("clock_alert", "")
        clock_notice = service.db.meta("clock_notice", "")
    daily = defaultdict(lambda: {"seconds": 0, "open_seconds": 0,
                                 "amount": Fraction(0), "open_amount": Fraction(0)})
    details = []
    names = {e["id"]: e["name"] for e in employees}
    for s in sessions:
        end = min(s["end_at"] if s["end_at"] is not None else now, stop, now)
        prefix = "open_" if s["end_at"] is None else ""
        for day, seconds in split_days(s["start_at"], max(s["start_at"], end), zone, hour):
            bucket = daily[(s["employee_id"], day)]
            bucket[prefix + "seconds"] += seconds
            bucket[prefix + "amount"] += Fraction(seconds * s["rate_cents"], 3600)
        if end > begin or (s["end_at"] is None and s["start_at"] < stop):
            notes = [s.get("review") or ""]
            if end - s["start_at"] > 86400:
                notes.append("24 saati aşıyor")
            details.append({**s, "name": names[s["employee_id"]],
                            "review": " · ".join(n for n in notes if n)})
    day_rows = []
    lifetime = defaultdict(int)
    period = defaultdict(lambda: {"seconds": 0, "open_seconds": 0, "cents": 0, "open_cents": 0})
    reviews = defaultdict(int)
    for d in details:
        if d.get("review"):
            reviews[d["employee_id"]] += 1
    for (eid, day), bucket in sorted(daily.items()):
        cents, open_cents = round_cents(bucket["amount"]), round_cents(bucket["open_amount"])
        lifetime[eid] += cents
        if first.isoformat() <= day <= last.isoformat():
            row = {"employee_id": eid, "name": names[eid], "day": day,
                   "seconds": bucket["seconds"], "open_seconds": bucket["open_seconds"],
                   "cents": cents, "open_cents": open_cents}
            day_rows.append(row)
            for key in period[eid]:
                period[eid][key] += row[key]
    total_paid, period_paid = defaultdict(int), defaultdict(int)
    payment_rows = []
    for p in payments:
        total_paid[p["employee_id"]] += p["amount_cents"]
        if p["paid_at"] >= begin:
            period_paid[p["employee_id"]] += p["amount_cents"]
            payment_rows.append({**p, "name": names[p["employee_id"]]})
    summary = []
    for e in employees:
        eid = e["id"]
        # Archived staff appear only when they still matter for this period.
        if not e["active"] and not (period[eid]["seconds"] or period[eid]["open_seconds"]
                                    or period_paid[eid] or lifetime[eid] - total_paid[eid]):
            continue
        summary.append({"employee_id": eid, "name": e["name"], **period[eid],
                        "paid_cents": period_paid[eid],
                        "period_net_cents": period[eid]["cents"] - period_paid[eid],
                        "carry_cents": lifetime[eid] - total_paid[eid]
                                       - period[eid]["cents"] + period_paid[eid],
                        "balance_cents": lifetime[eid] - total_paid[eid],
                        "review_count": reviews[eid]})
    return {"first": first.isoformat(), "last": last.isoformat(), "as_of": now,
            "timezone": str(zone), "day_start_hour": hour, "clock_alert": clock_alert, "clock_notice": clock_notice,
            "review_count": sum(1 for d in details if d.get("review")),
            "summary": summary,
            "days": day_rows, "sessions": details, "payments": payment_rows}


def safe_text(value):
    # Neutralize spreadsheet formulas, including leading control/whitespace.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def money(cents):
    return f"{cents / 100:.2f}"


def duration(seconds):
    hours, remain = divmod(seconds, 3600)
    minutes, seconds = divmod(remain, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def export_report(service, token, report, destination: Path):
    service.require_admin(token)
    destination = Path(destination)
    if destination.suffix.lower() not in (".csv", ".xlsx"):
        raise AppError("Dosya uzantısı .csv veya .xlsx olmalı.")
    summary_header = ["Personel No", "Ad Soyad", "Kapalı süre", "Ön hak ediş TL",
                      "Açık süre (geçici)", "Açık tahmini TL", "Dönem ödeme/avans TL",
                      "Dönem farkı TL", "Devreden TL", "Dönem sonu bakiye TL"]
    rows = [[r["employee_id"], r["name"], duration(r["seconds"]), money(r["cents"]),
             duration(r["open_seconds"]), money(r["open_cents"]), money(r["paid_cents"]),
             money(r["period_net_cents"]), money(r["carry_cents"]), money(r["balance_cents"])]
            for r in report["summary"]]
    fd, temporary = tempfile.mkstemp(dir=destination.parent, suffix=destination.suffix)
    os.close(fd)
    try:
        if destination.suffix.lower() == ".csv":
            with open(temporary, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f, delimiter=";")
                writer.writerow(["ÖN HAK EDİŞ — RESMÎ BORDRO DEĞİLDİR",
                                 report["first"], report["last"], report["timezone"]])
                writer.writerow(["Açık tutarlar bakiyeye dahil değildir.",
                                 service.display_time(report["as_of"]), report["clock_alert"],
                                 f"İnceleme bekleyen kayıt: {report.get('review_count', 0)}"])
                writer.writerow(summary_header)
                writer.writerows([
                    [Decimal(v) if i in (3, 5, 6, 7, 8, 9) else safe_text(v)
                     for i, v in enumerate(row)]
                    for row in rows
                ])
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            wb = Workbook()
            wb.remove(wb.active)

            def sheet(title, headers, data, money_columns=()):
                ws = wb.create_sheet(title)
                ws.append(headers)
                for row in data:
                    ws.append([safe_text(v) for v in row])
                    for column in money_columns:
                        cell = ws.cell(ws.max_row, column)
                        cell.value = float(row[column - 1])
                        cell.number_format = '#,##0.00'
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="173D3C")
                for i, cells in enumerate(ws.columns, 1):
                    ws.column_dimensions[get_column_letter(i)].width = min(
                        52, max(16, max(len(str(c.value or "")) for c in cells) + 2))
                return ws

            sheet("Özet", summary_header, rows, (4, 6, 7, 8, 9, 10))
            sheet("Günlük", ["Personel No", "Ad Soyad", "Tarih", "Kapalı süre", "Ön hak ediş TL",
                             "Açık süre", "Açık tahmini TL"],
                  [[r["employee_id"], r["name"], r["day"], duration(r["seconds"]), money(r["cents"]),
                    duration(r["open_seconds"]), money(r["open_cents"])] for r in report["days"]], (5, 7))
            sheet("Giriş çıkış", ["Kayıt", "Personel No", "Ad Soyad", "Giriş", "Çıkış",
                                  "Sabit saatlik TL", "Sürüm", "İnceleme"],
                  [[r["id"], r["employee_id"], r["name"], service.display_time(r["start_at"]),
                    service.display_time(r["end_at"]), money(r["rate_cents"]), r["revision"],
                    r["review"]] for r in report["sessions"]], (6,))
            sheet("Ödeme avans", ["Kayıt", "Personel No", "Ad Soyad", "Tarih", "Tür", "Tutar TL", "Not"],
                  [[r["id"], r["employee_id"], r["name"], service.display_time(r["paid_at"]),
                    r["kind"], money(r["amount_cents"]), r["note"]] for r in report["payments"]], (6,))
            sheet("Açıklama", ["Alan", "Değer"],
                  [["Belge", "Ön hak ediş; resmî bordro değildir."],
                   ["Dönem", report["first"] + " — " + report["last"]],
                   ["Üretim", service.display_time(report["as_of"])],
                   ["Saat dilimi", report["timezone"]],
                   ["İş günü", f"Her gün {report.get('day_start_hour', 0):02d}:00’da başlar; "
                               "bu saatten önceki çalışma önceki güne yazılır."], ["Saat uyarısı", report["clock_alert"]],
                   ["Saat bildirimi", report.get("clock_notice", "")],
                   ["İnceleme bekleyen kayıt", report.get("review_count", 0)],
                   ["Kural", "Saniye üzerinden hesaplanır; personel/gün başına kuruşa yuvarlanır."],
                   ["Açık kayıt", "Geçici tahmin; hak ediş ve bakiyeye dahil değildir."],
                   ["Ücret", "İşe girişte sabitlenir; değişiklik sonraki girişlere uygulanır."],
                   ["Bakiye", "Dönem sonuna kadar kapalı hak ediş eksi tüm ödeme ve avanslar."],
                   ["Kapsam", "Vergi, fazla mesai katsayısı, prim, izin ve kesinti hesabı içermez."]])
            wb.save(temporary)
            wb.close()
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


AUDIT_COLUMNS = ["id", "at", "actor", "action", "entity", "before_json", "after_json",
                 "reason", "previous_hash", "hash"]


def export_audit(service, token, destination):
    rows = service.audit_rows(token)
    destination = Path(destination)
    if destination.suffix.lower() != ".csv":
        raise AppError("İşlem geçmişi için .csv uzantısı kullanın.")
    fd, temporary = tempfile.mkstemp(dir=destination.parent, suffix=".csv")
    os.close(fd)
    try:
        with open(temporary, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(AUDIT_COLUMNS + ["Tarih (yerel)"])
            writer.writerows([[safe_text(row[c]) for c in AUDIT_COLUMNS] + [service.display_time(row["at"])]
                              for row in rows])
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
